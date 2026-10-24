# -*- coding: utf-8 -*-
"""复现 DEMO 大文件（12 条变更、35 个工序）的清单封面数据丢失问题。"""
import os
import sys
from datetime import datetime

BASE = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))  # 指导书工作台\
sys.path.insert(0, BASE)

import openpyxl
from openpyxl.styles import Font, Alignment, Border, Side

_TMP = os.environ.get("WI_TEST_TMP", "/tmp/wi_cover_tests")
os.makedirs(_TMP, exist_ok=True)
SRC = os.path.join(_TMP, "_repro_jpci_in.xlsx")
OUT = os.path.join(_TMP, "_repro_jpci_out.xlsx")

CHANGES = [
    ("2024-05-15", "Changed the working hours 更改工时"),
    ("2024-05-20", "The working hours of KA outer membrane welding disc were updated from 12S to 14S"),
    ("2024-05-20", "The working hours of KA inner membrane welding disc were updated from 11S to 14S"),
    ("2024-05-21", "Newly added automatic coil tube cutting machine 新增加卷管自动开料机"),
    ("2024-06-11", "The working hours for membrane grinding are changed from 1.5H/800pcs to 1.5H/2000pcs"),
    ("2024-06-21", "The testing equipment has been changed from the Muffle furnace to the vacuum oven"),
    ("2024-08-15", "The controller crimping process is updated to operate without a rubber sleeve"),
    ("2025-11-25", "增加 RA1000THT016A 热熔钻孔、RA1000THT016B 攻丝"),
    ("2026-09-11", "根据老板要求, 8GSD-22 渗漏检测工序因动作变更, 原工序工时 120 秒修改为 11 秒"),
    ("2026-09-14", "SGSD-22 渗漏检测工序动作变更"),
    ("2026-10-09", "KASD-1-1 工序中压铜鼻从 R2277-01 改为 F0234-1"),
    ("2026-10-20", "封面转换容量回归测试最后一条"),
]
AUTHORS = [
    ("2024-05-20", "自动"), ("2024-05-20", "自动"), ("2024-05-20", "Member"),
    ("2024-05-30", "清洗机"), ("2024-06-11", "P22压"), ("2024-07-12", "动力"),
    ("2024-09-02", "Diastat"), ("2026-08-11", "因球管"), ("2026-09-14", "SGSD"),
]
N_PROC = 35
FILE_NO = "DEMO-RDD-W-0054-2023(W23112812A-2)"

thin = Side(style="thin")
BORDER = Border(left=thin, right=thin, top=thin, bottom=thin)


def build():
    wb = openpyxl.Workbook()
    # ---------- 封面（老格式，DEMO 同构） ----------
    ws = wb.active
    ws.title = "清单"
    ws["A1"] = "Working Instruction Listing For stainless steel bulbs 作业指导书清单"
    ws["C2"] = "产品规格Specification"
    ws["D2"] = "KASD动力组件配线系列"
    ws["G2"] = "文件编号File NO"
    ws["I2"] = FILE_NO
    ws["G3"] = "版本Edition"
    ws["I3"] = "A15"
    ws["G4"] = "编制Made by"
    ws["I4"] = "邓嵘贤"
    ws["G5"] = "审核Checked by"
    ws["I5"] = "李春"
    ws["G6"] = "批准Approved by"
    ws["I6"] = "刘华峰"
    ws["G7"] = "发放日期Issued date"
    ws["I7"] = datetime(2026, 10, 9)
    # 变更表：左块 A/B 12 条，右块 D/E 作者备注 9 条（错行分布，同真实文件）
    ws["A4"] = "日期Date"
    ws["B4"] = "内容变更Change the content"
    for i, (d, c) in enumerate(CHANGES):
        r = 5 + i
        ws.cell(r, 1, datetime.strptime(d, "%Y-%m-%d"))
        ws.cell(r, 2, c)
    ws["D4"] = "日期Date"
    ws["E4"] = "作者Author"
    for i, (d, a) in enumerate(AUTHORS):
        r = 6 + i
        ws.cell(r, 4, datetime.strptime(d, "%Y-%m-%d"))
        ws.cell(r, 5, a)
    # ---------- 工序页（老格式） ----------
    for i in range(1, N_PROC + 1):
        p = wb.create_sheet(title=f"KASD-{i}")
        p["A1"] = "   作   业   指   导   书     Working Instruction"
        p["C2"] = "产品规格Specification"
        p["D2"] = "KASD动力组件配线系列"
        p["E2"] = "工时Working hour "
        p["F2"] = f"{i % 50 + 1}S"
        p["G2"] = "页码Page number"
        p["H2"] = str(i)
        p["C3"] = "产品名称Name"
        p["D3"] = "膜盒"
        p["E3"] = "文件编号File NO"
        p["G3"] = FILE_NO
        p["C4"] = "工位名称Testing part"
        p["D4"] = f"工序{i}名称"
        p["E4"] = "工序号Working procedure NO"
        p["G4"] = f"KASD-{i}"
        p["A5"] = "第一步：占位动作"
    wb.save(SRC)
    wb.close()
    print("输入文件已写出:", SRC, os.path.getsize(SRC), "字节")


def check():
    from excel_generator import normalize_file
    tmpl = os.path.join(BASE, "作业指导书版头.xlsx")
    wi = normalize_file(SRC, tmpl, OUT)
    print("规范化完成。解析到工序数:", len(wi.processes), " 变更记录数:", len(wi.header.revision_history))
    print("解析到的 file_no:", repr(wi.header.file_no))

    wb = openpyxl.load_workbook(OUT)
    ws = wb.worksheets[0]
    print("\n=== 输出封面 %s (max_row=%d) ===" % (ws.title, ws.max_row))
    print("I2 文件编号:", repr(ws["I2"].value))
    print("I3 版本:", repr(ws["I3"].value))
    print("I4 编制:", repr(ws["I4"].value))
    print("I7 发放日期:", repr(ws["I7"].value))
    print("\n-- 变更记录区（第 8 行起） --")
    for r in range(8, ws.max_row + 1):
        a, b = ws.cell(r, 1).value, ws.cell(r, 2).value
        if a is not None or b is not None:
            print(f"  行{r}: A={str(a)[:24]!r} B={str(b)[:50]!r}")
    print("\n-- 工序清单区 --")
    n = 0
    for r in range(15, ws.max_row + 1):
        a = ws.cell(r, 1).value
        if a and str(a).strip() and "Working" not in str(a):
            n += 1
    print("  列出的工序行数:", n, "（输入 %d 个）" % N_PROC)
    # J44 SUM 公式位置
    for r in range(40, ws.max_row + 1):
        v = ws.cell(r, 10).value
        if isinstance(v, str) and v.startswith("=SUM"):
            print("  SUM 公式位于 J%d: %s" % (r, v))
    wb.close()
    print("\n输出文件:", OUT, os.path.getsize(OUT), "字节")


if __name__ == "__main__":
    build()
    check()
