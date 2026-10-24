# -*- coding: utf-8 -*-
"""封面扩容边界用例：生成路径超容、恰好满容、零记录、无封面文件。"""
import os
import sys

BASE = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
sys.path.insert(0, BASE)

import openpyxl
from excel_generator import ExcelGenerator, normalize_file
from wi_models import WorkInstruction, Header, Process, RevisionRecord

TMPL = os.path.join(BASE, "作业指导书版头.xlsx")
_TMP = os.environ.get("WI_TEST_TMP", "/tmp/wi_cover_tests")
os.makedirs(_TMP, exist_ok=True)
OUT = os.path.join(_TMP, "_edge_out.xlsx")


def make_wi(n_proc, n_rev):
    wi = WorkInstruction()
    wi.header = Header(file_no="EDGE-TEST-001", edition="B2",
                       made_by="测试员", checked_by="审核员",
                       approved_by="批准员", issued_date="2026/10/24",
                       product_spec="测试规格")
    for i in range(n_rev):
        wi.header.revision_history.append(
            RevisionRecord(date=f"2026-10-{i + 1:02d}", content=f"第{i + 1}条变更内容", author=""))
    for i in range(n_proc):
        p = Process()
        p.process_no = f"EDGE-{i + 1}"
        p.process_name_cn = f"边界工序{i + 1}"
        p.work_hours = i + 1
        wi.add_process(p)
    return wi


def cover_summary(path):
    import re
    wb = openpyxl.load_workbook(path)
    ws = wb.worksheets[0]
    n_rev = n_proc = 0
    date_re = re.compile(r"^\d{4}[-/年.]\d{1,2}[-/月.]\d{1,2}$")
    for r in range(9, ws.max_row + 1):
        a = ws.cell(r, 1).value
        if a is None:
            continue
        s = str(a).strip()
        if not s or "Working" in s or "Date" in s:
            continue
        if date_re.match(s):
            n_rev += 1
        else:
            n_proc += 1
    sum_row = None
    for r in range(ws.max_row, 9, -1):
        v = ws.cell(r, 10).value
        if isinstance(v, str) and v.startswith("=SUM"):
            sum_row = (r, v)
            break
    wb.close()
    return n_rev, n_proc, sum_row


def case(name, n_proc, n_rev):
    wi = make_wi(n_proc, n_rev)
    gen = ExcelGenerator(TMPL)
    gen.generate(wi, OUT)
    n_rev_out, n_proc_out, sum_row = cover_summary(OUT)
    ok = (n_rev_out == n_rev and n_proc_out == n_proc)
    print(f"[{'PASS' if ok else 'FAIL'}] {name}: 变更 {n_rev}→{n_rev_out}, "
          f"工序 {n_proc}→{n_proc_out}, SUM {sum_row}")
    return ok


results = []
results.append(case("生成路径·35 工序/12 变更（双超容）", 35, 12))
results.append(case("生成路径·恰满容 29 工序/4 变更", 29, 4))
results.append(case("生成路径·零变更零工序", 0, 0))
results.append(case("生成路径·仅工序超容 40/2", 40, 2))
results.append(case("生成路径·仅变更超容 5/9", 5, 9))

# 无封面文件（老格式但无清单页）：结构不应被改动
src = os.path.join(_TMP, "_edge_nocover_in.xlsx")
wb = openpyxl.Workbook()
ws = wb.active
ws.title = "工序页"
ws["A1"] = "作业指导书 Working Instruction"
ws["G4"] = "NC-1"
wb.save(src)
wb.close()
out2 = os.path.join(_TMP, "_edge_nocover_out.xlsx")
normalize_file(src, TMPL, out2)
wb = openpyxl.load_workbook(out2)
names = wb.sheetnames
wb.close()
ok = names == ["工序页"]
print(f"[{'PASS' if ok else 'FAIL'}] 无封面文件表数不变: {names}")
results.append(ok)

print()
print("全部通过 ✅" if all(results) else "存在失败项 ❌")
